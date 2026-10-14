import io
import logging
from typing import Optional

logger = logging.getLogger("extractor")

def extract_text_from_bytes(content: bytes, filename: str, ext: str) -> str:
    ext = ext.lower().lstrip(".")
    text = ""
    
    try:
        if ext == "pdf":
            # PyMuPDF extraction
            try:
                import fitz
                doc = fitz.open(stream=content, filetype="pdf")
                pages_text = []
                for i in range(min(len(doc), 50)):  # up to 50 pages
                    pages_text.append(doc[i].get_text())
                text = "\n".join(pages_text)
            except Exception as e:
                logger.warning(f"PyMuPDF failed on {filename}: {e}, falling back to pypdf")
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(content))
                text = "\n".join([page.extract_text() or "" for page in reader.pages[:50]])

        elif ext in ["docx", "doc"]:
            import docx
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

        elif ext in ["xlsx", "xls"]:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            sheet_texts = []
            for sheetname in wb.sheetnames[:5]:
                ws = wb[sheetname]
                for row in ws.iter_rows(max_row=100, values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        sheet_texts.append(" | ".join(row_vals))
            text = "\n".join(sheet_texts)

        elif ext in ["txt", "md", "csv", "json", "py", "ts", "js", "html", "css", "xml", "yaml", "yml", "log"]:
            try:
                text = content.decode("utf-8", errors="ignore")
            except Exception:
                text = content.decode("latin-1", errors="ignore")

        else:
            # Non-text binary (image, video, audio, archive)
            text = ""

    except Exception as e:
        logger.error(f"Text extraction failed for {filename} ({ext}): {e}")
        text = ""

    # Sanitize and limit length to reasonable boundary (e.g. 60,000 chars)
    return text.strip()[:60000]
